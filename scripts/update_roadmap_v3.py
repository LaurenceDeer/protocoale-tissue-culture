import openpyxl
import time
import sys

def update_roadmap():
    sys.stdout.reconfigure(encoding='utf-8')
    filename = 'ROADMAP-HUB.xlsx'
    max_retries = 3
    for attempt in range(1, max_retries + 1):
        try:
            wb = openpyxl.load_workbook(filename)
            
            # Update Sheet: 2. Calendar 8 săpt
            ws2 = wb['2. Calendar 8 săpt']
            updated_cal = 0
            for r in range(2, ws2.max_row + 1):
                task_id = str(ws2.cell(row=r, column=2).value).strip()
                if task_id == 'S7':
                    ws2.cell(row=r, column=1, value='Da')
                    ws2.cell(row=r, column=8, value='2026-07-06')
                    updated_cal += 1

            wb.save(filename)
            print(f"SUCCES: S-au actualizat {updated_cal} randuri in '2. Calendar 8 sapt' (S7 bifat ca Da).")
            return True
        except PermissionError:
            print(f"AVERTISMENT (Incercarea {attempt}/{max_retries}): Fisierul '{filename}' este deschis in Excel. Inchideti Excel si reincercati...")
            if attempt < max_retries:
                time.sleep(2)
            else:
                print(f"EROARE: Nu s-a putut salva '{filename}' deoarece este blocat de un alt proces (Excel).")
                return False

if __name__ == '__main__':
    update_roadmap()
