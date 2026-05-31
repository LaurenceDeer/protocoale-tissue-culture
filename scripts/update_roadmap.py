import openpyxl
import time
import os
import sys

def update_roadmap():
    sys.stdout.reconfigure(encoding='utf-8')
    filename = 'ROADMAP-HUB.xlsx'
    max_retries = 3
    for attempt in range(1, max_retries + 1):
        try:
            wb = openpyxl.load_workbook(filename)
            
            # 1. Update Sheet: 1. Faze & Taskuri
            ws1 = wb['1. Faze & Taskuri']
            updated_faze = 0
            for r in range(2, ws1.max_row + 1):
                task_id = str(ws1.cell(row=r, column=2).value).strip()
                if task_id in ['2.1', '2.2', '2.4', '2.5']:
                    ws1.cell(row=r, column=1, value='Da')
                    ws1.cell(row=r, column=8, value='2026-05-27')
                    updated_faze += 1

            # 2. Update Sheet: 2. Calendar 8 săpt
            ws2 = wb['2. Calendar 8 săpt']
            updated_cal = 0
            for r in range(2, ws2.max_row + 1):
                task_id = str(ws2.cell(row=r, column=2).value).strip()
                if task_id in ['S4', 'S5']:
                    ws2.cell(row=r, column=1, value='Da')
                    ws2.cell(row=r, column=8, value='2026-05-27')
                    updated_cal += 1

            wb.save(filename)
            print(f"SUCCES: S-au actualizat {updated_faze} task-uri in '1. Faze & Taskuri' si {updated_cal} randuri in '2. Calendar 8 sapt'.")
            return True
        except PermissionError:
            print(f"AVERTISMENT (Incercarea {attempt}/{max_retries}): Fisierul '{filename}' este deschis in Excel. Inchideti Excel si reincercati...")
            if attempt < max_retries:
                time.sleep(2)
            else:
                print(f"EROARE: Nu s-a putut salva '{filename}' deoarece este blocat de un alt proces (Excel). Va rugam sa rulati manual scriptul: python scripts/update_roadmap.py dupa ce inchideti fisierul.")
                return False

if __name__ == '__main__':
    update_roadmap()
