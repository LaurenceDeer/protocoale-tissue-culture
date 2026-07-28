import openpyxl
import time
import sys

def update_roadmap():
    sys.stdout.reconfigure(encoding='utf-8')
    filename = 'ROADMAP-HUB.xlsx'
    max_retries = 3
    for attempt in range(1, max_retries + 1):
        try:
            wb = openpyxl.load_workbook(filename)
            
            ws = wb['4. Amânate & Viitor']
            updated_count = 0
            for r in range(2, ws.max_row + 1):
                task_id = str(ws.cell(row=r, column=2).value).strip()
                if task_id == 'A-11':
                    ws.cell(row=r, column=1, value='Da')
                    ws.cell(row=r, column=8, value='2026-07-28')
                    updated_count += 1

            wb.save(filename)
            print(f"SUCCES: S-a bifat task-ul A-11 ca 'Da' în '4. Amânate & Viitor'.")
            return True
        except PermissionError:
            print(f"AVERTISMENT (Incercarea {attempt}/{max_retries}): Fisierul '{filename}' este deschis in Excel. Inchideti Excel si reincercati...")
            if attempt < max_retries:
                time.sleep(2)
            else:
                print(f"EROARE: Nu s-a putut salva '{filename}' deoarece este blocat de un alt proces (Excel).")
                return False

if __name__ == '__main__':
    update_roadmap()
